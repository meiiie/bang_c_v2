"""Only query-selected workbooks may contribute data to prompts or programs."""
import json
from pathlib import Path
import shutil
import subprocess

import openpyxl
import pytest

from solver.llm import MockClient
from solver.pipeline import Pipeline, input_hash
from solver.program import run_program


@pytest.fixture
def dataset(tmp_path):
    root = tmp_path / 'input'
    books = root / 'xlsx_files'
    books.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active.append(['Name', 'Score'])
    book.active.append(['Alice', 7])
    book.save(books / 'scores.xlsx')
    (root / 'queries.json').write_text(json.dumps([
        {'id': 'q1', 'file_name': 'scores.xlsx', 'query': 'What is the score for Alice?'}
    ]))
    return root


def test_dump_ignores_answer_files(dataset, tmp_path, monkeypatch):
    from main import main
    clean = tmp_path / 'clean'
    shutil.copytree(dataset, clean)
    for name in ('answers.json', 'sample_submission.csv', 'gold.csv', 'gold.json',
                 'private_submission.csv', 'secret_answer.txt'):
        (dataset / name).write_text('LEAKED_GOLD_SENTINEL')
    # Even unused workbooks must never be read by provenance hashing.
    (dataset / 'xlsx_files' / 'gold.xlsx').write_text('LEAKED_GOLD_SENTINEL')
    real_open = Path.open
    def check_input_read(path, mode='r', *args, **kwargs):
        if 'r' in mode:
            for root in (clean, dataset):
                if path.is_relative_to(root):
                    assert path in (root / 'queries.json', root / 'xlsx_files/scores.xlsx')
        return real_open(path, mode, *args, **kwargs)
    with monkeypatch.context() as guard:
        guard.setattr(Path, 'open', check_input_read)
        for label, root in [('clean', clean), ('extra', dataset)]:
            assert main(['--input', str(root), '--output', str(tmp_path / f'{label}.csv'),
                         '--mode', 'dump', '--dump-dir', str(tmp_path / label / 'prompts'),
                         '--no-checkpoint']) == 0
    prompt = (clean / 'prompts/q1.txt').read_text()
    assert 'Alice' in prompt
    assert 'LEAKED_GOLD_SENTINEL' not in prompt
    assert prompt == (tmp_path / 'extra/prompts/q1.txt').read_text()
    assert input_hash(dataset) == input_hash(clean)


@pytest.mark.parametrize('name', [
    '../answers.json', '../outside.xlsx', '/tmp/outside.xlsx',
    r'..\outside.xlsx', r'C:\outside.xlsx', 'answers.json',
    'sample_submission.csv', 'scores.xlsx/../answers.json', 'scores.xlsx\x00.csv',
])
def test_pipeline_rejects_invalid_workbooks(dataset, tmp_path, name):
    query = json.loads((dataset / 'queries.json').read_text())[0]
    pipeline = Pipeline(MockClient(), log_path=tmp_path / 'log.jsonl')
    pipeline.solve(dict(query, file_name=name), dataset, mode='dump', dump_dir=tmp_path / 'prompts')
    record = json.loads(pipeline.log_path.read_text())
    assert record['error'].startswith('ValueError: Invalid workbook filename')
    assert not (tmp_path / 'prompts/q1.txt').exists()


def test_pipeline_rejects_workbook_symlink_escape(dataset, tmp_path):
    outside = tmp_path / 'outside.xlsx'
    shutil.copy2(dataset / 'xlsx_files/scores.xlsx', outside)
    (dataset / 'xlsx_files/escape.xlsx').symlink_to(outside)
    test_pipeline_rejects_invalid_workbooks(dataset, tmp_path, 'escape.xlsx')


def test_checkpoint_does_not_hash_escaped_workbook(dataset, tmp_path, monkeypatch):
    outside = tmp_path / 'outside.xlsx'
    shutil.copy2(dataset / 'xlsx_files/scores.xlsx', outside)
    (dataset / 'xlsx_files/escape.xlsx').symlink_to(outside)
    query = json.loads((dataset / 'queries.json').read_text())[0]
    query['file_name'] = 'escape.xlsx'
    real_open = Path.open
    def check_read(path, *args, **kwargs):
        assert path.resolve() != outside
        return real_open(path, *args, **kwargs)
    monkeypatch.setattr(Path, 'open', check_read)
    pipeline = Pipeline(MockClient(), log_path=tmp_path / 'log.jsonl')
    pipeline.solve_all([query], dataset, output_path=tmp_path / 'output.csv')
    records = [json.loads(line) for line in pipeline.log_path.read_text().splitlines()]
    assert records
    assert all(record['error'].startswith('ValueError: Invalid workbook filename') for record in records)
    assert not pipeline.checkpoint_keys


@pytest.mark.parametrize('expression', [
    "open('../answers.json').read()",
    "pd.read_csv('sample_submission.csv')",
    "pd.io.common.get_handle('../answers.json', 'r').handle.read()",
    "pd.io.parsers.TextFileReader('sample_submission.csv').read()",
])
def test_program_cannot_read_answer_files(dataset, monkeypatch, expression):
    (dataset / 'answers.json').write_text('LEAKED_GOLD_SENTINEL')
    (dataset / 'xlsx_files/sample_submission.csv').write_text('answer\nLEAKED_GOLD_SENTINEL\n')
    monkeypatch.chdir(dataset / 'xlsx_files')
    result = run_program('answer = ' + expression, dataset / 'xlsx_files/scores.xlsx')
    assert 'error' in result
    assert 'LEAKED_GOLD_SENTINEL' not in str(result)
    assert any(word in result['error'] for word in ('Forbidden', 'PermissionError'))


def test_program_blocks_absolute_reads_and_writes(dataset):
    secret = dataset / 'sample_submission.csv'
    secret.write_text('answer\nLEAKED_GOLD_SENTINEL\n')
    workbook = dataset / 'xlsx_files/scores.xlsx'
    for target, mode in [(secret, 'r'), (workbook, 'w'), (secret, 'w')]:
        result = run_program(f'answer = pd.io.common.get_handle({str(target)!r}, {mode!r}).handle.read()', workbook)
        assert 'PermissionError' in result.get('error', '')
        assert 'LEAKED_GOLD_SENTINEL' not in str(result)
    assert secret.read_text() == 'answer\nLEAKED_GOLD_SENTINEL\n'
    assert openpyxl.load_workbook(workbook).active['B2'].value == 7


def test_program_allows_current_workbook_and_sheet_helper(dataset):
    workbook = dataset / 'xlsx_files/scores.xlsx'
    code = (f'header = pd.io.common.get_handle({str(workbook)!r}, "rb", is_text=False).handle.read(2)\n'
            'answer = [list(header), sheet(tables, "Sheet")["Score"].sum()]')
    assert run_program(code, workbook) == {'answer': [[80, 75], 7]}


def test_program_blocks_process_escape(dataset):
    secret = dataset / 'answers.json'
    secret.write_text('LEAKED_GOLD_SENTINEL')
    result = run_program(f'answer = pd.io.common.os.system("cat {secret}")',
                         dataset / 'xlsx_files/scores.xlsx')
    assert 'PermissionError' in result.get('error', '')
    assert 'LEAKED_GOLD_SENTINEL' not in str(result)


def test_program_cannot_replace_allowed_workbook(dataset):
    secret = dataset / 'answers.json'
    secret.write_text('LEAKED_GOLD_SENTINEL')
    workbook = dataset / 'xlsx_files/scores.xlsx'
    code = (f'pd.io.common.os.rename({str(secret)!r}, {str(workbook)!r})\n'
            f'answer = pd.io.common.get_handle({str(workbook)!r}, "r").handle.read()')
    result = run_program(code, workbook)
    assert 'PermissionError' in result.get('error', '')
    assert 'LEAKED_GOLD_SENTINEL' not in str(result)
    assert secret.exists()
    assert openpyxl.load_workbook(workbook).active['B2'].value == 7


@pytest.mark.parametrize('target', [
    'pd.io.common.os.path.realpath', 'pd.io.common.os.fspath',
    "pd.io.common.os.sys.modules['builtins'].isinstance",
    "pd.io.common.os.walk('/tmp').gi_frame.f_globals['fspath']",
])
def test_program_cannot_patch_file_guard(dataset, target):
    workbook = dataset / 'xlsx_files/scores.xlsx'
    secret = dataset / 'answers.json'
    secret.write_text('LEAKED_GOLD_SENTINEL')
    code = (f'{target} = lambda *args: {str(workbook)!r}\n'
            f'answer = pd.io.common.get_handle({str(secret)!r}, "r").handle.read()')
    result = run_program(code, workbook)
    assert 'Forbidden' in result.get('error', '')
    assert 'LEAKED_GOLD_SENTINEL' not in str(result)


@pytest.mark.parametrize('hook', ['settrace', 'setprofile'])
def test_program_cannot_inspect_guard_frames(dataset, hook):
    code = f'answer = pd.io.common.os.sys.{hook}(lambda *args: None)'
    result = run_program(code, dataset / 'xlsx_files/scores.xlsx')
    assert 'PermissionError' in result.get('error', '')


def test_program_cannot_retrieve_forbidden_builtins_via_pickle(dataset):
    workbook = dataset / 'xlsx_files/scores.xlsx'
    code = (f'handle = pd.io.common.get_handle({str(workbook)!r}, "rb", is_text=False).handle\n'
            'answer = pd.compat.pickle_compat.Unpickler(handle).find_class("builtins", "setattr")(pd.io.common.os.path, "realpath", lambda p: "spoofed")')
    result = run_program(code, workbook)
    assert 'PermissionError' in result.get('error', '')


def test_program_uses_disposable_empty_cwd(dataset, monkeypatch):
    real_run = subprocess.run
    working_dirs = []
    def capture(*args, **kwargs):
        cwd = Path(kwargs['cwd'])
        assert cwd.is_dir()
        assert list(cwd.iterdir()) == []
        working_dirs.append(cwd)
        return real_run(*args, **kwargs)
    monkeypatch.setattr('solver.program.subprocess.run', capture)
    for _ in range(2):
        assert run_program('answer = len(tables[0])', dataset / 'xlsx_files/scores.xlsx') == {'answer': 1}
    assert len(set(working_dirs)) == 2
    assert all(not path.exists() for path in working_dirs)
