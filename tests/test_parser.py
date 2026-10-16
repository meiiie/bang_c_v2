from datetime import datetime
from pathlib import Path
import openpyxl
import pytest
from solver.parser import parse_workbook, clean_text, is_data, is_missing_marker, normalize_value


def load(data_dir, n):
    return parse_workbook(data_dir/'xlsx_files'/f'p4_t{n:02}.xlsx')


def test_hours_and_agriculture(data_dir):
    t = load(data_dir, 8)[0]
    f = t.to_dataframe()
    c = next(c for c in f if 'Persons at work > Agriculture' in c)
    assert f.loc['5 to 14 hours', c] == 134
    assert t.units == ''


def test_unpaid_family_workers(data_dir):
    f = load(data_dir, 6)[0].to_dataframe()
    c = next(c for c in f if 'Agriculture and related industries > Unpaid family workers' in c)
    assert f.loc['Total, 16 years and over', c] == 35
    assert any('Wage and salary workers > Private industries > Total' in c for c in f)


def test_industry_headers(data_dir):
    t = load(data_dir, 5)[0]
    assert '2023 > Construction' in t.header_paths
    assert '2023 > Manufacturing' in t.header_paths
    assert t.units == 'In thousands'
    assert len(t.rows) == 77
    assert all('Footnotes' not in r.label_path and 'HOUSEHOLD' not in r.label_path for r in t.rows)


def test_split_and_formula(data_dir):
    ts = load(data_dir, 2)
    assert len(ts) == 2
    assert ts[0].header_paths[0] == 'Fiscal years > 2020'
    assert ts[1].header_paths[0] == 'Calendar years > 2020'
    assert ts[1].header_paths[-1] == 'September > 2024'
    assert all(t.units == 'Billion dollars' for t in ts)
    assert ts[0].rows[-1].label_path == 'Trade balance'


def test_sections(data_dir):
    f = load(data_dir, 4)[0].to_dataframe()
    assert 'Women > Unemployment rate' in f.index
    assert 'TOTAL > Unemployment rate' in f.index
    t = load(data_dir, 10)[0]
    assert any(r.label_path == 'MARITAL STATUS > Women, 16 years and over > Widowed, divorced, or separated' for r in t.rows)


def test_indent(data_dir):
    t = load(data_dir, 3)[0]
    paths = {r.label_path for r in t.rows}
    assert any(len(path.split(' > ')) >= 4 for path in paths)
    for path in paths:
        levels = path.split(' > ')
        assert all(' > '.join(levels[:depth]) in paths for depth in range(1, len(levels)))
    assert 'Sep. 2024- Oct. 2024' in t.header_paths[5]


def test_people(data_dir):
    t = load(data_dir, 1)[0]
    assert len(t.rows) == 20
    assert t.header_paths[:3] == ['Name > First', 'Name > Middle', 'Name > Last']
    assert isinstance(t.rows[0].values[3], datetime)
    assert t.rows[0].label_path == '1'
    assert '2005-01-03T00:00:00' in t.to_text(10000)
    assert len(t.to_text(150)) <= 150


def test_text_cleanup():
    assert clean_text(' Note-\nbook ') == 'Notebook'
    assert clean_text('Green-\nhouse') == 'Greenhouse'
    assert clean_text('unincor-\nporated') == 'unincorporated'
    assert clean_text('Total(2)') == 'Total'
    assert clean_text('Net inventory 3/') == 'Net inventory'
    assert clean_text('Total(2)', False) == 'Total(2)'


def test_reference_chain_and_cycle(tmp_path):
    w = openpyxl.Workbook(); s = w.active
    s.append(['Label','Value']); s.append(['one', 42]); s.append(['two','=B2']); s.append(['three','=B3'])
    s.append(['cycle','=B6']); s.append(['cycle2','=B5'])
    path=tmp_path/'refs.xlsx'; w.save(path)
    t=parse_workbook(path)[0]
    assert t.to_dataframe().loc['three', 'Value'] == 42


def test_age_at_event_column_optional(data_dir):
    from solver.config import load_config
    config = load_config()
    config["enrichers"] = ["age_at_event"]
    from solver.parser import parse_workbook
    t = parse_workbook(data_dir/'xlsx_files/p4_t01.xlsx', config)[0]
    assert t.header_paths[-1] == 'Age at Admission (days)'
    assert t.rows[0].values[-1] == 7345


def test_no_default_age_enrichment(data_dir):
    assert not any(h.startswith('Age at') for h in load(data_dir, 1)[0].header_paths)


@pytest.mark.parametrize('marker', [
    '-', '–', '—', '−', '(-8-)', '(−8−)', '(–8–)', '(—8—)',
    ' ( - 8 - ) ', 'NA', 'n.a.', 'N/A',
])
def test_missing_markers(marker):
    assert is_missing_marker(marker)
    assert normalize_value(marker) is None
    assert is_data(['Total all sizes', 3, marker, marker, marker])


@pytest.mark.parametrize('value,expected', [
    ('12.5%', 12.5), ('(1,234)', -1234), ('(8)', -8),
    ('(-8)', 8), ('−8', -8), ('(12.5)%', -12.5),
])
def test_numeric_normalization_preserved(value, expected):
    assert not is_missing_marker(value)
    assert normalize_value(value) == expected


def test_missing_size_class_preserves_industry_and_headers():
    path = Path(__file__).parent / 'fixtures' / 'missing_size_classes.xlsx'
    table = parse_workbook(path)[0]
    row = next(r for r in table.rows if r.label_path.endswith('250 - 999'))
    assert row.label_path == 'Fixture industry > 250 - 999'
    assert row.values[0] == 25.3
    assert table.header_paths == [
        'Average incidence rates > (mean)',
        'First quartile > (1st quartile)',
        'Median > (median)',
        'Third quartile > (3rd quartile)',
    ]


@pytest.mark.parametrize('label', [
    '   1 - 10', '11 - 49', '50 – 249', '250 − 999', '1,000+',
])
def test_empty_size_classes_preserve_section(tmp_path, label):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in [
        ['Label', 'Mean'], ['Fixture industry(2)', None],
        [label, None], ['   Observed class', 25.3],
    ]:
        sheet.append(row)
    path = tmp_path / 'empty_size_class.xlsx'
    workbook.save(path)
    workbook.close()
    table = parse_workbook(path)[0]
    assert table.rows[0].label_path == 'Fixture industry > Observed class'
    assert table.rows[0].values == [25.3]


def test_missing_observation_preserves_section(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in [
        ['Label', 'Mean'], ['Fixture industry', None],
        ['Unreported category', '(-8-)'], ['Reported category', 4.5],
    ]:
        sheet.append(row)
    path = tmp_path / 'missing_observation.xlsx'
    workbook.save(path)
    workbook.close()
    table = parse_workbook(path)[0]
    assert table.rows[0].label_path == 'Fixture industry > Reported category'
